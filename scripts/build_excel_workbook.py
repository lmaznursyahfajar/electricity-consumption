"""
build_excel_workbook.py
========================
Membangun `data/dataset_konsumsi_listrik.xlsx` — workbook Excel berisi
sampel dataset yang diformat rapi untuk keperluan eksplorasi cepat &
dokumentasi (bukan pengganti CSV lengkap, yang tetap jadi sumber data
utama untuk aplikasi karena jauh lebih cepat dibaca pandas).

Jalankan dari root proyek:
    python scripts/build_excel_workbook.py
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src import config  # noqa: E402

SAMPLE_ROWS = 5000
FONT_NAME = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="0F62FE")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="0F62FE")
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="6B7280")
BODY_FONT = Font(name=FONT_NAME, size=10)
THIN_BORDER = Border(*(Side(style="thin", color="D1D5DB"),) * 4)


def style_header_row(ws, row_idx, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def autosize_columns(ws, df, start_col=1, max_width=38):
    for i, col in enumerate(df.columns, start=start_col):
        width = min(max(len(str(col)), df[col].astype(str).str.len().quantile(0.9) if len(df) else 10) + 3,
                    max_width)
        ws.column_dimensions[get_column_letter(i)].width = width


def write_dataframe(ws, df: pd.DataFrame, start_row: int = 1, table_name: str | None = None):
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=start_row, column=j, value=str(col))
    style_header_row(ws, start_row, len(df.columns))

    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    autosize_columns(ws, df)

    if table_name:
        last_col_letter = get_column_letter(len(df.columns))
        last_row = start_row + len(df)
        ref = f"A{start_row}:{last_col_letter}{last_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
        )
        ws.add_table(table)


def build_workbook():
    fact = pd.read_csv(config.FACT_CSV_PATH)
    dim_pelanggan = pd.read_csv(config.DIM_PELANGGAN_PATH)
    dim_wilayah = pd.read_csv(config.DIM_WILAYAH_PATH)

    sample = fact.sample(n=min(SAMPLE_ROWS, len(fact)), random_state=42).sort_values(
        ["customer_id", "date"]
    ).reset_index(drop=True)

    wb = Workbook()

    # --- Sheet 1: Ringkasan (dashboard formula) ---
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Dataset Konsumsi Listrik — Ringkasan"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Dataset sintetis untuk latihan analisis pola konsumsi listrik & segmentasi "
                "pelanggan dengan Self-Organizing Map (SOM).")
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A1:D1")
    ws.merge_cells("A2:D2")

    labels = [
        ("Jumlah pelanggan (tabel Dim_Pelanggan)", "=COUNTA(Dim_Pelanggan!A2:A100000)"),
        ("Jumlah wilayah (tabel Dim_Wilayah)", "=COUNTA(Dim_Wilayah!A2:A100000)"),
        ("Jumlah baris pada sampel fakta", "=COUNTA(Sampel_Fakta!A2:A100000)"),
        ("Total konsumsi pada sampel (kWh)", "=SUM(Sampel_Fakta!C2:C100000)"),
        ("Rata-rata konsumsi pada sampel (kWh)", "=AVERAGE(Sampel_Fakta!C2:C100000)"),
        ("Rata-rata suhu pada sampel (°C)", "=AVERAGE(Sampel_Fakta!D2:D100000)"),
        ("Jumlah baris anomali pada sampel", "=SUM(Sampel_Fakta!I2:I100000)"),
    ]
    start = 4
    ws.cell(row=start, column=1, value="Metrik").font = HEADER_FONT
    ws.cell(row=start, column=2, value="Nilai").font = HEADER_FONT
    for r, c in [(start, 1), (start, 2)]:
        ws.cell(row=r, column=c).fill = HEADER_FILL
        ws.cell(row=r, column=c).border = THIN_BORDER
    for i, (label, formula) in enumerate(labels, start=start + 1):
        ws.cell(row=i, column=1, value=label).font = BODY_FONT
        cell = ws.cell(row=i, column=2, value=formula)
        cell.font = BODY_FONT
        ws.cell(row=i, column=1).border = THIN_BORDER
        cell.border = THIN_BORDER
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22
    ws.cell(row=start + len(labels) + 2, column=1,
            value="Catatan: metrik konsumsi di atas dihitung dari sheet 'Sampel_Fakta' "
                  f"({SAMPLE_ROWS:,} baris acak), bukan keseluruhan dataset. "
                  "Dataset lengkap (seluruh baris) tersedia pada file CSV terpisah — lihat sheet 'Petunjuk'."
            ).font = NOTE_FONT
    ws.merge_cells(start_row=start + len(labels) + 2, start_column=1,
                    end_row=start + len(labels) + 2, end_column=6)

    # --- Sheet 2: Sampel Fakta Konsumsi ---
    ws2 = wb.create_sheet("Sampel_Fakta")
    write_dataframe(ws2, sample, table_name="TabelSampelFakta")

    # --- Sheet 3: Dim Pelanggan ---
    ws3 = wb.create_sheet("Dim_Pelanggan")
    write_dataframe(ws3, dim_pelanggan, table_name="TabelDimPelanggan")

    # --- Sheet 4: Dim Wilayah ---
    ws4 = wb.create_sheet("Dim_Wilayah")
    write_dataframe(ws4, dim_wilayah, table_name="TabelDimWilayah")

    # --- Sheet 5: Kamus Data ---
    kamus = pd.DataFrame([
        ("fact_konsumsi_harian", "customer_id", "Teks", "ID unik pelanggan, kunci relasi ke Dim_Pelanggan"),
        ("fact_konsumsi_harian", "date", "Tanggal", "Tanggal pencatatan konsumsi (harian)"),
        ("fact_konsumsi_harian", "consumption_kwh", "Numerik", "Konsumsi listrik harian dalam kWh (dapat kosong/NaN)"),
        ("fact_konsumsi_harian", "temperature_c", "Numerik", "Suhu rata-rata harian sintetis (°C)"),
        ("fact_konsumsi_harian", "peak_load_kwh", "Numerik", "Estimasi beban puncak harian (kWh)"),
        ("fact_konsumsi_harian", "load_factor", "Numerik (0-1)", "Rasio beban rata-rata terhadap beban puncak"),
        ("fact_konsumsi_harian", "is_weekend", "0/1", "1 jika tanggal jatuh pada akhir pekan"),
        ("fact_konsumsi_harian", "is_holiday", "0/1", "1 jika tanggal adalah hari libur nasional (simulasi)"),
        ("fact_konsumsi_harian", "is_anomaly", "0/1", "1 jika baris disuntik sebagai anomali konsumsi"),
        ("dim_pelanggan", "customer_id", "Teks", "ID unik pelanggan (primary key)"),
        ("dim_pelanggan", "customer_type", "Kategori", "Segmen pelanggan (Rumah Tangga, Bisnis, Industri, dst.)"),
        ("dim_pelanggan", "tariff_class", "Kategori", "Golongan tarif yang disederhanakan (R1/900VA, B2, I2, dst.)"),
        ("dim_pelanggan", "daya_terpasang_va", "Numerik", "Daya listrik terpasang dalam VA"),
        ("dim_pelanggan", "tarif_rp_per_kwh", "Numerik", "Tarif ilustratif Rp/kWh — BUKAN tarif resmi PLN"),
        ("dim_pelanggan", "region", "Kategori", "Kota/wilayah pelanggan, kunci relasi ke Dim_Wilayah"),
        ("dim_pelanggan", "base_consumption_kwh", "Numerik", "Baseline konsumsi harian sebelum faktor musiman/acak"),
        ("dim_pelanggan", "tanggal_pasang", "Tanggal", "Tanggal instalasi meter listrik pelanggan"),
        ("dim_wilayah", "region", "Kategori", "Nama kota/wilayah (primary key)"),
        ("dim_wilayah", "provinsi", "Kategori", "Provinsi tempat wilayah berada"),
        ("dim_wilayah", "suhu_baseline_c", "Numerik", "Suhu rata-rata tahunan baseline wilayah (°C)"),
        ("dim_wilayah", "amplitudo_musiman", "Numerik", "Faktor pengali variasi musiman suhu & konsumsi"),
    ], columns=["Tabel", "Kolom", "Tipe Data", "Deskripsi"])
    ws5 = wb.create_sheet("Kamus_Data")
    write_dataframe(ws5, kamus, table_name="TabelKamusData")

    # --- Sheet 6: Petunjuk ---
    ws6 = wb.create_sheet("Petunjuk")
    ws6["A1"] = "Petunjuk Penggunaan Dataset"
    ws6["A1"].font = TITLE_FONT
    petunjuk_lines = [
        "",
        "1. Dataset LENGKAP (seluruh baris, bukan sampel) tersedia dalam format CSV pada folder 'data/':",
        "   - fact_konsumsi_harian.csv  (data transaksional harian per pelanggan)",
        "   - dim_pelanggan.csv         (data induk pelanggan)",
        "   - dim_wilayah.csv           (data induk wilayah)",
        "",
        "2. Workbook Excel ini hanya memuat SAMPEL dari fact_konsumsi_harian (agar tetap ringan",
        "   dibuka di Excel), plus tabel dimensi LENGKAP dan kamus data.",
        "",
        "3. Untuk menjalankan aplikasi analisis (Streamlit), gunakan file CSV di folder 'data/' —",
        "   aplikasi akan memuatnya secara otomatis. Lihat README.md untuk cara menjalankan aplikasi.",
        "",
        "4. Ingin membuat ulang dataset dengan parameter berbeda (jumlah pelanggan, rentang tanggal,",
        "   dsb.)? Ubah nilai di src/config.py lalu jalankan:  python -m src.data_generator",
        "",
        "5. Seluruh angka tarif (Rp/kWh) bersifat ILUSTRATIF untuk simulasi dan BUKAN tarif resmi PLN.",
    ]
    for i, line in enumerate(petunjuk_lines, start=2):
        ws6.cell(row=i, column=1, value=line).font = BODY_FONT
    ws6.column_dimensions["A"].width = 95

    wb.move_sheet("Ringkasan", offset=-5)
    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(config.EXCEL_WORKBOOK_PATH)
    print(f"Workbook tersimpan di {config.EXCEL_WORKBOOK_PATH}")


if __name__ == "__main__":
    build_workbook()
